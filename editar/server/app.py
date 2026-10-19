import random
import re
from flask import Flask, request, jsonify, send_from_directory, render_template, redirect, session, url_for, send_file
import pandas as pd
import os
import openpyxl
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from datetime import datetime
import numpy as np
import json

# Setup Flask and Flask-Login
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
app = Flask(__name__, template_folder=os.path.join(ROOT_DIR, 'templates'),
            static_folder=os.path.join(ROOT_DIR, 'website'), static_url_path='/_s')
app.config["APPLICATION_ROOT"] = '/acl'
app.secret_key = os.environ.get('SECRET_KEY') or os.urandom(32)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'  # Name of the login view for Flask-Login to redirect to

EXCEL_PATH = os.environ.get('EXCEL_PATH', os.path.join(ROOT_DIR, 'database.xlsx'))
EDIT_PASSWORD = os.environ.get('EDIT_PASSWORD')
# Dummy user data
users = {}

# User class for Flask-Login
class User(UserMixin):
    def __init__(self, user_id):
        self.id = user_id

@login_manager.user_loader
def load_user(user_id):
    if user_id in users:
        return User(user_id)
    return None


def normalize_text(value):
    if pd.isna(value):
        return ''
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def is_missing_text(value):
    return normalize_text(value).lower() in ('', 'nan', 'none')


def parse_date_bounds(value):
    text = normalize_text(value)
    if not text:
        return None

    text = re.sub(r'^(\d+)\.0$', r'\1', text)
    text = re.sub(r'^(\d+)\.0\s*-\s*(\d+)\.0$', r'\1-\2', text)

    if re.match(r'^\d+$', text):
        year = int(text)
        return year, year

    match = re.match(r'^(\d+)\s*-\s*(\d+)$', text)
    if not match:
        return None

    start, end = int(match.group(1)), int(match.group(2))
    if start > end:
        start, end = end, start
    return start, end


def date_contains_year(value, year):
    bounds = parse_date_bounds(value)
    if not bounds:
        return False
    start, end = bounds
    return start <= year <= end


def build_year_counts(df, id_col='id', date_col='date'):
    year_counts = {}
    seen_pairs = set()

    for row in df[[id_col, date_col]].itertuples(index=False):
        record_id, date_value = row
        bounds = parse_date_bounds(date_value)
        if not bounds:
            continue

        start, end = bounds
        for year in range(start, end + 1):
            pair = (record_id, year)
            if pair in seen_pairs:
                continue
            seen_pairs.add(pair)
            year_counts[year] = year_counts.get(year, 0) + 1

    return dict(sorted(year_counts.items()))


@app.route('/images', methods=['GET'])
def get_images():
    df = load_excel_data()
    df.sort_values(by='id', inplace=True)  # Ensure correct order
    df.fillna("", inplace=True)
    data = df.to_dict(orient='records')
    return jsonify(data)


@app.route('/verify-password', methods=['POST'])
@app.route('/acl/verify-password', methods=['POST'])
def verify_password():
    data = request.get_json()
    password = data.get('password', '')
    
    if EDIT_PASSWORD and password == EDIT_PASSWORD:
        return jsonify({'success': True}), 200
    else:
        return jsonify({'success': False, 'message': 'Password incorreta'}), 401


@app.route('/images/<int:image_id>', methods=['POST'])
def update_image(image_id):
    df = load_excel_data()
    image = df.loc[df['id'] == image_id]
    if image.empty:
        return "Image not found", 404
    info = request.json.get('info', '')
    df.loc[df['id'] == image_id, 'info'] = info
    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet = wb.worksheets[0]
    headers = {cell.value: cell.column for cell in sheet[1]}
    info_column = headers.get('info', sheet.max_column + 1)
    sheet.cell(1, info_column, 'info')
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == image_id:
            sheet.cell(row[0].row, info_column, info)
            break
    save_catalog_workbook(wb)
    return "Success", 200


# Redirect /acl to /acl/
@app.route('/acl', methods=['GET'])
def missing_bar():
    return redirect('/acl/', code=302)


@app.route('/acl/', methods=['GET'])
def acl_root():
    return home()


@app.route('/', methods=['GET'])
def home():
    # Load the Excel data
    df = load_excel_data()

    for col in ['classification', 'lingua', 'author', 'editor', 'titulo', 'local', 'local(atual)', 'estado']:
        df[col] = df[col].apply(normalize_text)
    df['date'] = df['date'].apply(normalize_text)

    # Get filters from query parameters
    filters = {
        'estado': request.args.get('estado', ''),
        'classification': request.args.get('classification', ''),
        'lingua': request.args.get('language', ''),
        'author': request.args.get('author', ''),
        'editor': request.args.get('editor', ''),
        'local': request.args.get('local', ''),
        'local(atual)': request.args.get('local_atual', ''),
        'date': request.args.get('date', ''),
    }

    # Apply filters to the dataframe
    filtered_df = df.copy()
    for key, value in filters.items():
        if value:
            if key == 'date':
                if value == '_sem valor':
                    filtered_df = filtered_df[filtered_df['date'].apply(is_missing_text)]
                else:
                    try:
                        year = int(value)
                        filtered_df = filtered_df[filtered_df['date'].apply(lambda date_value: date_contains_year(date_value, year))]
                    except ValueError:
                        escaped_value = re.escape(value)
                        filtered_df = filtered_df[filtered_df['date'].str.contains(escaped_value, case=False, na=False)]
            elif value == '_sem valor':
                filtered_df = filtered_df[filtered_df[key].apply(is_missing_text)]
            else:
                escaped_value = re.escape(value)
                filtered_df = filtered_df[filtered_df[key].astype(str).str.contains(escaped_value, case=False, na=False)]

    # Generate lists and counts for filtering options
    classification_list = sorted([value for value in df['classification'].unique().tolist() if not is_missing_text(value)])
    author_list = sorted([value for value in df['author'].unique().tolist() if not is_missing_text(value)])
    editor_list = sorted([value for value in df['editor'].unique().tolist() if not is_missing_text(value)])
    title_list = sorted([value for value in df['titulo'].unique().tolist() if not is_missing_text(value)])
    local_list = sorted([value for value in df['local'].unique().tolist() if not is_missing_text(value)])
    language_list = sorted([value for value in df['lingua'].unique().tolist() if not is_missing_text(value)])
    local_atual_list = sorted([value for value in df['local(atual)'].unique().tolist() if not is_missing_text(value)])
    date_counts = build_year_counts(filtered_df)
    date_list = list(date_counts.keys())

    filtered_unique_df = filtered_df.drop_duplicates(subset=['id'], keep='first').reset_index(drop=True)

    # Counts for each filtering option
    classification_counts = filtered_unique_df['classification'].value_counts().to_dict()
    author_counts = filtered_unique_df['author'].value_counts().to_dict()
    editor_counts = filtered_unique_df['editor'].value_counts().to_dict()
    title_counts = filtered_unique_df['titulo'].value_counts().to_dict()
    local_counts = filtered_unique_df['local'].value_counts().to_dict()
    language_counts = filtered_unique_df['lingua'].value_counts().to_dict()
    local_atual_counts = filtered_unique_df['local(atual)'].value_counts().to_dict()
    estado_counts = filtered_unique_df['estado'].value_counts().to_dict()

    classification_counts['_sem valor'] = int(df['classification'].apply(is_missing_text).sum())
    author_counts['_sem valor'] = int(df['author'].apply(is_missing_text).sum())
    editor_counts['_sem valor'] = int(df['editor'].apply(is_missing_text).sum())
    title_counts['_sem valor'] = int(df['titulo'].apply(is_missing_text).sum())
    local_counts['_sem valor'] = int(df['local'].apply(is_missing_text).sum())
    language_counts['_sem valor'] = int(df['lingua'].apply(is_missing_text).sum())
    local_atual_counts['_sem valor'] = int(df['local(atual)'].apply(is_missing_text).sum())
    date_counts['_sem valor'] = int(df['date'].apply(is_missing_text).sum())
    estado_counts['_sem valor'] = int(df['estado'].apply(is_missing_text).sum())

    # Render the template with all the data
    return render_template('indices.html',
                           classifications=classification_list,
                           authors=author_list,
                           editors=editor_list,
                           estado=sorted([value for value in df['estado'].unique().tolist() if not is_missing_text(value)]),
                           titles=title_list,
                           locals=local_list,
                           locals_atual=local_atual_list,
                           language=language_list,
                           date=date_list,
                           classification_counts=classification_counts,
                           author_counts=author_counts,
                           editor_counts=editor_counts,
                           title_counts=title_counts,
                           local_counts=local_counts,
                           local_atual_counts=local_atual_counts,
                           language_counts=language_counts,
                           date_counts=date_counts,
                           estado_counts=estado_counts,
                           filters=filters  # Pass filters to the template
                          )





@app.route('/download-excel')
def download_excel():
    # Path to your Excel file
    excel_path = os.path.abspath(EXCEL_PATH)

    # Get current date and time
    current_datetime = datetime.now().strftime('%Y%m%d%H%M')  # Format: YYYYMMDDHHMM

    # Generate the filename with the current timestamp
    filename = f"database-{current_datetime}.xlsx"

    # Serve the Excel file for download with the new filename
    return send_file(excel_path, as_attachment=True, download_name=filename)


@app.route('/download-catalogo')
def download_catalogo():
    catalogo_path = 'catalogo.pdf'

    return send_file(catalogo_path, as_attachment=True)

@app.route('/<path:path>', methods=['GET'])
def serve_file_in_dir(path):
    if not os.path.isfile(os.path.join(app.static_folder, path)):
        path = os.path.join(path, 'index.html')
    return send_from_directory(app.static_folder, path)

@app.route('/submit-description', methods=['POST'])
def submit_description():
    data = request.get_json()

    if not data or 'id' not in data or 'description' not in data:
        return jsonify({'message': 'Missing id or description'}), 400

    id = int(data.get('id'))
    description = data.get('description')

    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet2 = wb['Sheet2'] if 'Sheet2' in wb.sheetnames else wb.worksheets[0]
    description_column = 3 if sheet2.title == 'Sheet2' else next(
        cell.column for cell in sheet2[1] if cell.value == 'description')

    row_found = False
    for row in sheet2.iter_rows(min_row=2):
        if row[0].value == id:
            sheet2.cell(row[0].row, description_column, description)
            row_found = True
            break

    if not row_found:
        return jsonify({'message': 'ID not found'}), 404

    save_catalog_workbook(wb)

    return jsonify({'message': 'Description received'}), 200

@app.route('/image-details/<int:image_id>', methods=['GET'])
def show_image_details(image_id):
    df = load_excel_data()

    df['estado'] = df['estado'].fillna('').astype(str)

    # Get the image data for the given ID
    image_data = df[df['id'] == image_id].to_dict('records')

    # If no image is found, return a 404 error
    if not image_data:
        return "Image not found", 404

    # Get the first (and only) record
    image_data = image_data[0]

    # Replace `nan`, 0, and None values with empty strings
    for key, value in image_data.items():
        if pd.isna(value) or value == 0 or value is None:
            image_data[key] = ""

    # Get the path and other metadata
    image_path = image_data['path'].replace('\\', '/')
    date = image_data.get('date', '')
    author = image_data.get('author', '')
    editor = image_data.get('editor', '')
    url = image_data.get('URL', '')

    path_number_match = re.search(r'/(\d+)_', image_path)
    path_number = str(int(path_number_match.group(1))) if path_number_match else "N/A"

    # Load chat messages from the Excel file
    chat_wb = openpyxl.load_workbook(EXCEL_PATH)
    chat_sheet = chat_wb['Chat'] if 'Chat' in chat_wb.sheetnames else None
    messages = []
    for row in chat_sheet.iter_rows(min_row=2, values_only=True) if chat_sheet else []:
        if row[0] == image_id:
            if len(row) < 5:
                row = list(row) + [0]
            else:
                row = row[:5]
            messages.append({'name': row[1], 'message': row[2], 'timestamp': row[3]})

    classification_list = df['classification'].unique().tolist()

    # Render the details page, passing in the image data and messages
    return render_template('image_details.html', classifications=classification_list, image={
        'id': image_id,
        'path': image_path,
        'path_number': path_number,  # Include the extracted path number
        'info': image_data['Bibliographic Description'],
        'date': date,
        'author': author,
        'editor': editor,
        'editor_name': image_data.get('editor_name', ''),
        'tradutor': image_data.get('tradutor', ''),
        'url': url,
        'estado': image_data['estado'],
        'classification': image_data.get('classification', ''),
        'titulo': image_data.get('titulo', ''),
        'local(atual)': image_data.get('local(atual)', ''),
        'local': image_data.get('local', ''),
        'tomos': image_data.get('tomos', ''),
        'notas': image_data.get('notas', ''),
        'carreira': image_data.get('carreira', ''),
        'lingua': image_data.get('lingua', ''),
        'registo': image_data.get('registo', ''),
        'messages': messages
    })




@app.route('/chat/<int:image_id>', methods=['GET', 'POST'])
def chat(image_id):
    if request.method == 'POST':
        name = request.json.get('name', '')
        if not name:
            name = "User" + str(random.randint(1000, 9999))
        message = request.json.get('message', '')

        if message:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            if 'Chat' not in wb.sheetnames:
                chat_sheet = wb.create_sheet('Chat')
                chat_sheet.append(['image_id', 'name', 'message', 'timestamp', 'viewed'])
            else:
                chat_sheet = wb['Chat']
            chat_sheet.append([image_id, name, message, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), 0])
            save_catalog_workbook(wb)

    wb = openpyxl.load_workbook(EXCEL_PATH)
    chat_sheet = wb['Chat'] if 'Chat' in wb.sheetnames else None
    messages = []
    for row in chat_sheet.iter_rows(min_row=2, values_only=True) if chat_sheet else []:
        if row[0] == image_id:
            messages.append({'name': row[1], 'message': row[2], 'timestamp': row[3]})

    return jsonify(messages)

def load_excel_data():
    df = pd.read_excel(EXCEL_PATH, sheet_name=0)
    df = df.rename(columns={'Impressor': 'editor', 'impressor': 'editor', 'Editor': 'editor'})
    return df


def save_catalog_workbook(workbook):
    """Keep the public column name without changing the internal API field."""
    for cell in workbook.worksheets[0][1]:
        if str(cell.value).casefold() in ('editor', 'impressor'):
            cell.value = 'Impressor'
            break
    workbook.save(EXCEL_PATH)
    workbook.close()

@app.route('/indices')
def indices():
    df = load_excel_data()
    classifications = df['classification'].unique().tolist()
    authors = df['author'].unique().tolist()
    editors = df['editor'].unique().tolist()
    titles = df['titulo'].unique().tolist()
    locals = df['local'].unique().tolist()
    locals_atual = df['local(atual)'].unique().tolist()
    date = df['date'].unique().tolist()
    return render_template('indices.html', classifications=classifications, authors=authors,editors=editors ,titles=titles, locals=locals,locals_atual=locals_atual,date = date)

@app.route('/search', methods=['POST'])
def search():
    data = request.get_json()
    page = data.get('page', 1)
    per_page = data.get('per_page', 15)

    # Load the Excel sheet
    df = load_excel_data()

    # Capture the initial count of empty values before conversion
    initial_counts = {
        'classification': int(df['classification'].apply(is_missing_text).sum()),
        'lingua': int(df['lingua'].apply(is_missing_text).sum()),
        'author': int(df['author'].apply(is_missing_text).sum()),
        'editor': int(df['editor'].apply(is_missing_text).sum()),
        'titulo': int(df['titulo'].apply(is_missing_text).sum()),
        'local': int(df['local'].apply(is_missing_text).sum()),
        'local(atual)': int(df['local(atual)'].apply(is_missing_text).sum()),
        'date': int(df['date'].apply(is_missing_text).sum()),
        'estado': int(df['estado'].apply(is_missing_text).sum())
    }

    # Convert necessary columns to strings to avoid errors and preserve
    # the canonical date interval representation.
    for col in ['classification', 'lingua', 'author', 'editor', 'titulo', 'local', 'local(atual)', 'date', 'estado']:
        df[col] = df[col].apply(normalize_text)

    # Filters for each category
    filters = {
        'classification': data.get('classification'),
        'lingua': data.get('language'),
        'author': data.get('author'),
        'editor': data.get('editor'),
        'titulo': data.get('titulo'),
        'local': data.get('local'),
        'local(atual)': data.get('local_atual'),
        'date': data.get('date'),
        'estado': data.get('estado'),
        'search_word': data.get('search_word')
    }

    # Apply filters to the dataframe
    filtered_df = df.copy()
    for key, value in filters.items():
        if value:
            if key == 'date':
                if value == "_sem valor":
                    filtered_df = filtered_df[filtered_df['date'].apply(is_missing_text)]
                else:
                    try:
                        year = int(value)
                        filtered_df = filtered_df[filtered_df['date'].apply(lambda date_value: date_contains_year(date_value, year))]
                    except ValueError:
                        escaped_value = re.escape(value)
                        filtered_df = filtered_df[filtered_df['date'].str.contains(escaped_value, case=False, na=False)]
            elif value == "_sem valor":  # Handle "sem valor" specifically
                filtered_df = filtered_df[filtered_df[key].apply(is_missing_text)]
            elif key == 'search_word':
                filtered_df = filtered_df[filtered_df.apply(lambda row: row.astype(str).str.contains(value, case=False).any(), axis=1)]
            elif key == 'classification':  # Add special handling for classification
                # Use exact matching for classification
                filtered_df = filtered_df[filtered_df[key].str.strip() == value]
            else:
                escaped_value = re.escape(value)
                filtered_df = filtered_df[filtered_df[key].astype(str).str.contains(escaped_value, case=False, na=False)]

    # Unique records for the visible results and total count.
    filtered_unique_df = filtered_df.drop_duplicates(subset=['id'], keep='first').reset_index(drop=True)

    # Calculate value counts for each filter category based on the filtered dataframe
    classification_counts = filtered_unique_df['classification'].value_counts().to_dict()
    language_counts = filtered_unique_df['lingua'].value_counts().to_dict()
    author_counts = filtered_unique_df['author'].value_counts().to_dict()
    editor_counts = filtered_unique_df['editor'].value_counts().to_dict()
    title_counts = filtered_unique_df['titulo'].value_counts().to_dict()
    local_counts = filtered_unique_df['local'].value_counts().to_dict()
    local_atual_counts = filtered_unique_df['local(atual)'].value_counts().to_dict()
    date_counts = build_year_counts(filtered_unique_df)
    estado_counts = filtered_unique_df['estado'].value_counts().to_dict()

    # Include "sem valor" counts using initial counts (before filtering)
    classification_counts["_sem valor"] = initial_counts['classification']
    language_counts["_sem valor"] = initial_counts['lingua']
    author_counts["_sem valor"] = initial_counts['author']
    editor_counts["_sem valor"] = initial_counts['editor']
    title_counts["_sem valor"] = initial_counts['titulo']
    local_counts["_sem valor"] = initial_counts['local']
    local_atual_counts["_sem valor"] = initial_counts['local(atual)']
    date_counts["_sem valor"] = initial_counts['date']
    estado_counts["_sem valor"] = initial_counts['estado']

    # Total results for the current filtered dataframe
    total_results = len(filtered_unique_df)

    # Pagination: Select rows for the current page
    df = filtered_unique_df.iloc[(page - 1) * per_page:page * per_page]

    # Ensure all columns are converted to standard Python data types
    df = df.astype(object).where(pd.notnull(df), None)

    # Prepare results in dictionary format for JSON response
    results = df.replace({np.nan: ''}).to_dict(orient='records')

    # Convert counts dictionaries to standard Python types
    classification_counts = {str(k): int(v) for k, v in classification_counts.items()}
    language_counts = {str(k): int(v) for k, v in language_counts.items()}
    author_counts = {str(k): int(v) for k, v in author_counts.items()}
    editor_counts = {str(k): int(v) for k, v in editor_counts.items()}
    title_counts = {str(k): int(v) for k, v in title_counts.items()}
    local_counts = {str(k): int(v) for k, v in local_counts.items()}
    local_atual_counts = {str(k): int(v) for k, v in local_atual_counts.items()}
    date_counts = {str(k): int(v) for k, v in date_counts.items()}
    estado_counts = {str(k): int(v) for k, v in estado_counts.items()}

    # Send the updated filter counts and results to the frontend
    return jsonify({
        'current_page': page,
        'results': results,
        'total_results': total_results,
        'filter_counts': {
            'classification_counts': classification_counts,
            'language_counts': language_counts,
            'author_counts': author_counts,
            'editor_counts': editor_counts,
            'title_counts': title_counts,
            'local_counts': local_counts,
            'local_atual_counts': local_atual_counts,
            'date_counts': date_counts,
            'estado_counts': estado_counts
        }
    })




@app.route('/update-image-new-info', methods=['POST'])
def update_image_new_info():
    data = request.get_json()
    image_id = data.get('id')

    if not image_id:
        return jsonify({'message': 'ID not provided'}), 400

    try:
        # Open the workbook once and work with both sheets sequentially
        wb = openpyxl.load_workbook(EXCEL_PATH)
        sheet1 = wb.worksheets[0]

        # Update data in Sheet1
        for row in sheet1.iter_rows(min_row=2):
            if row[0].value == image_id:
                row[3].value = data.get('info', row[2].value)  # Bibliographic Description
                row[5].value = data.get('classification', row[4].value)  # Classification
                row[6].value = data.get('author', '').replace('\n', '\n')  # Preserve newlines
                row[7].value = data.get('titulo', row[6].value)  # Title
                row[8].value = data.get('date', row[7].value)  # Date
                row[9].value = data.get('local', row[9].value)  # Local
                row[10].value = data.get('tomos', row[10].value)  # Tomos
                row[11].value = data.get('url', '').replace('\n', '\n')  # Preserve newlines
                row[12].value = data.get('lingua', row[12].value)  # Lingua
                row[13].value = data.get('carreira', row[13].value)  # Career
                row[14].value = data.get('notas', row[14].value)  # Notas
                row[15].value = data.get('registo', row[15].value)  # Registo
                row[16].value = data.get('editor', '').replace('\n', '\n')  # Preserve newlines
                row[17].value = data.get('local_atual', row[17].value)  # Local(atual)
                row[18].value = data.get('estado', row[18].value)  # Estado
                row[19].value = data.get('editor_name', '').replace('\n', '\n')  # Editor
                row[20].value = data.get('tradutor', '').replace('\n', '\n')  # Tradutor
                break

        # If a new chat message is present, append it to the Chat sheet
        new_message = data.get('chat_message', '')
        new_name = data.get('chat_name', f"User{random.randint(1000, 9999)}")

        if new_message:
            if 'Chat' not in wb.sheetnames:
                chat_sheet = wb.create_sheet('Chat')
                chat_sheet.append(['image_id', 'name', 'message', 'timestamp', 'viewed'])
            else:
                chat_sheet = wb['Chat']
            chat_sheet.append([image_id, new_name, new_message, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), 0])

        # Save the workbook after all operations are done
        save_catalog_workbook(wb)

    except Exception as e:
        return jsonify({'message': f"Error updating the Excel file: {e}"}), 500

    return jsonify({'message': 'Changes and chat saved successfully'}), 200



@app.route('/catalog-search', methods=['POST'])
def catalog_search():
    data = request.get_json()
    page = int(data.get('page', 1))  # Ensure the page number is an integer

    # Adjust the page to start from the appropriate Excel pages (based on your catalog numbering)
    adjusted_page = page + 2

    per_page = data.get('per_page', 15)  # Get how many images per page, default is 15

    # Load data from Excel
    df = load_excel_data()

    # Handle missing or NaN values in the 'page' column
    df = df.dropna(subset=['page'])  # Drop rows where 'page' is NaN

    # Convert the 'page' column to integers to match the incoming page number
    df['page'] = df['page'].astype(int)

    # Filter the images for the given page and the next page (use adjusted_page and adjusted_page + 1)
    df_filtered_left = df[df['page'] == adjusted_page]  # Current page images
    df_filtered_right = df[df['page'] == adjusted_page + 1]  # Next page images

    total_images_left = len(df_filtered_left)
    total_images_right = len(df_filtered_right)

    results_left = df_filtered_left.replace({np.nan: ''}).to_dict(orient='records')
    results_right = df_filtered_right.replace({np.nan: ''}).to_dict(orient='records')

    return jsonify({
        'current_page': page,
        'results_left': results_left,
        'results_right': results_right,
        'total_images_left': total_images_left,
        'total_images_right': total_images_right
    })




if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000,debug=True)
