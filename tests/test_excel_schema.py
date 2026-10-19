import importlib.util
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from server import app as public


ROOT = Path(__file__).resolve().parents[1]
editorial = None
if (ROOT / 'editar/server/app.py').exists() and importlib.util.find_spec('flask_login'):
    spec = importlib.util.spec_from_file_location('editorial_app', ROOT / 'editar/server/app.py')
    editorial = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(editorial)


@unittest.skipIf(editorial is None, 'Install requirements-dev.txt to test both applications')
class ExcelSchemaTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / 'database.xlsx'
        wb = openpyxl.load_workbook(ROOT / 'database.xlsx')
        for sheet in list(wb)[1:]:
            wb.remove(sheet)
        wb.save(self.path)
        wb.close()
        self.public_path = patch.object(public, 'EXCEL_PATH', str(self.path))
        self.editorial_path = patch.object(editorial, 'EXCEL_PATH', str(self.path))
        self.public_path.start()
        self.editorial_path.start()

    def tearDown(self):
        self.editorial_path.stop()
        self.public_path.stop()
        self.directory.cleanup()

    def test_both_apps_accept_new_and_legacy_header(self):
        for header in ['Impressor', 'editor', 'Editor']:
            wb = openpyxl.load_workbook(self.path)
            wb.worksheets[0]['Q1'] = header
            wb.save(self.path)
            wb.close()
            self.assertEqual(len(public.load_data()), 1869)
            self.assertIn('editor', public.load_data())
            self.assertIn('editor', editorial.load_excel_data())

    def test_public_single_sheet_routes_and_export(self):
        client = public.app.test_client()
        for route in ['/catalogo', '/obra/1834', '/download-excel']:
            response = client.get(route)
            self.assertEqual(response.status_code, 200, route)
            if route == '/download-excel':
                wb = openpyxl.load_workbook(BytesIO(response.data))
                self.assertEqual(wb.worksheets[0]['Q1'].value, 'Impressor')
                self.assertEqual(wb.sheetnames, ['Sheet1'])
                wb.close()
            response.close()

    def test_editorial_single_sheet_routes(self):
        client = editorial.app.test_client()
        for route in ['/', '/images', '/image-details/1834', '/chat/1834', '/download-excel']:
            response = client.get(route)
            self.assertEqual(response.status_code, 200, route)
            response.close()

    def test_save_without_chat_keeps_impressor(self):
        response = editorial.app.test_client().post('/update-image-new-info', json={
            'id': 1834, 'editor': 'Impressor de teste',
        })
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(self.path)
        self.assertEqual(wb.worksheets[0]['Q1'].value, 'Impressor')
        row = next(row for row in wb.worksheets[0].iter_rows(min_row=2) if row[0].value == 1834)
        self.assertEqual(row[16].value, 'Impressor de teste')
        self.assertEqual(wb.sheetnames, ['Sheet1'])
        wb.close()

    def test_comment_creates_chat_only_when_needed(self):
        client = editorial.app.test_client()
        response = client.post('/chat/1834', json={'name': 'Teste', 'message': 'Comentário'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json[0]['message'], 'Comentário')
        wb = openpyxl.load_workbook(self.path)
        self.assertEqual(wb.worksheets[0]['Q1'].value, 'Impressor')
        self.assertIn('Chat', wb.sheetnames)
        wb.close()

    def test_description_without_sheet2(self):
        response = editorial.app.test_client().post('/submit-description', json={
            'id': 1834, 'description': 'Descrição de teste',
        })
        self.assertEqual(response.status_code, 200)

    def test_legacy_info_update_preserves_auxiliary_sheets(self):
        wb = openpyxl.load_workbook(self.path)
        wb.worksheets[0]['Q1'] = 'editor'
        wb.create_sheet('Chat').append(['image_id', 'name', 'message', 'timestamp', 'viewed'])
        wb.save(self.path)
        wb.close()
        response = editorial.app.test_client().post('/images/1834', json={'info': 'Teste'})
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(self.path)
        self.assertEqual(wb.worksheets[0]['Q1'].value, 'Impressor')
        self.assertIn('Chat', wb.sheetnames)
        wb.close()

    def test_missing_edit_password_fails_closed(self):
        with patch.object(editorial, 'EDIT_PASSWORD', None):
            response = editorial.app.test_client().post('/verify-password', json={'password': ''})
        self.assertEqual(response.status_code, 401)
